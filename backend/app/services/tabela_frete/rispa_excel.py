"""Leitor da tabela comercial RISPA com tarifas por faixa de CEP."""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

from app.services.tabela_frete.analise import AnaliseDocumentoError


FORMATO = "uf_zona_peso_v1"


def _numero(valor: object) -> float | None:
    if valor in (None, "", "N/A", "NA", "-"):
        return None
    texto = str(valor).strip().replace("R$", "").replace(" ", "")
    if "," in texto:
        texto = texto.replace(".", "").replace(",", ".")
    try:
        return float(texto)
    except ValueError as exc:
        raise AnaliseDocumentoError(f"Valor numérico inválido na tabela RISPA: {valor}") from exc


def _cep(valor: object) -> str | None:
    digitos = re.sub(r"\D", "", str(valor or ""))
    return digitos.zfill(8) if digitos else None


def _cabecalho(aba) -> list[str]:
    return [str(aba.cell(1, coluna).value or "").strip().upper() for coluna in range(1, 7)]


def extrair_rispa_excel(caminho: Path) -> dict:
    try:
        workbook = load_workbook(caminho, data_only=True, read_only=True)
    except Exception as exc:
        raise AnaliseDocumentoError(f"Não foi possível ler o Excel: {exc}") from exc

    try:
        aba = workbook["TABELA FAIXA CEP"]
        if _cabecalho(aba) != ["UF", "CIDADE", "SIGLA", "CEPI", "CEPF", "PRAZO D UTIL"]:
            raise AnaliseDocumentoError("Cabeçalho da malha RISPA não reconhecido")

        tarifas: dict[tuple[str, str], dict] = {}
        proximos_indices: dict[tuple[str, str], int] = {}
        mapeamento: dict[str, list[dict]] = {}
        prazos: dict[str, int] = {}
        localidades_sem_cep = 0

        for linha in range(2, aba.max_row + 1):
            uf = str(aba.cell(linha, 1).value or "").strip().upper()
            cidade = str(aba.cell(linha, 2).value or "").strip()
            zona = str(aba.cell(linha, 3).value or "").strip().upper()
            if not uf and not cidade and not zona:
                continue
            if len(uf) != 2 or not cidade or not zona:
                raise AnaliseDocumentoError(f"UF, cidade ou sigla inválida na linha {linha}")

            faixas = [
                {"ate_kg": peso, "valor": valor}
                for peso, valor in zip(
                    (20, 30, 50, 70, 100),
                    (_numero(aba.cell(linha, coluna).value) for coluna in range(8, 13)),
                )
                if valor is not None
            ]
            if not faixas:
                raise AnaliseDocumentoError(f"Nenhuma tarifa encontrada na linha {linha}")

            tarifa = {
                "uf": uf,
                "uf_nome": uf,
                "zona": zona,
                "faixas_peso": faixas,
                "excedente_por_kg_acima_100": _numero(aba.cell(linha, 14).value) or 0.0,
                "gris_percentual": _numero(aba.cell(linha, 15).value) or 0.0,
                "ad_valorem_percentual": 0.0,
                "pedagio_por_fracao_100kg": _numero(aba.cell(linha, 17).value) or 0.0,
                "tas_por_cte": _numero(aba.cell(linha, 16).value) or 0.0,
                "trt": _numero(aba.cell(linha, 18).value),
            }
            chave = (uf, zona)
            anterior = tarifas.get(chave)
            if anterior is not None and anterior != tarifa:
                indice = proximos_indices.get(chave, 2)
                zona = f"{zona}#{indice}"
                tarifa["zona"] = zona
                chave = (uf, zona)
                proximos_indices[(uf, str(aba.cell(linha, 3).value or "").strip().upper())] = indice + 1
            tarifas[chave] = tarifa

            cep_inicio = _cep(aba.cell(linha, 4).value)
            cep_fim = _cep(aba.cell(linha, 5).value)
            if not cep_inicio or not cep_fim:
                localidades_sem_cep += 1
            try:
                prazo = int(float(aba.cell(linha, 6).value))
            except (TypeError, ValueError) as exc:
                raise AnaliseDocumentoError(f"Prazo inválido na linha {linha}") from exc
            item = {
                "cidade": cidade,
                "uf": uf,
                "zona": zona,
                "prazo_dias": prazo,
                "cep_inicio": cep_inicio,
                "cep_fim": cep_fim,
                "tda": _numero(aba.cell(linha, 7).value) or 0.0,
                "trt": tarifa["trt"],
                "bloqueio_entrega": False,
                "bloqueio_coleta": False,
                "bloqueio_ambos": False,
            }
            mapeamento.setdefault(f"{uf}|{zona}", []).append(item)
            prazos[f"{uf}|{cidade}"] = prazo

        if not tarifas:
            raise AnaliseDocumentoError("Nenhuma tarifa RISPA encontrada")
        return {
            "formato": FORMATO,
            "origem": {"descricao": "Guarulhos - SP", "cidade": "Guarulhos", "uf": "SP"},
            "tipo_calculo": "EXCEDENTE_ACIMA_100KG",
            "fator_cubagem": 300.0,
            "tarifas_por_zona": list(tarifas.values()),
            "mapeamento_zonas": mapeamento,
            "prazos_entrega": prazos,
            "regras_gerais": {
                "pedagio": "por fração de 100 kg",
                "icms": "conforme legislação em vigor; não informado numericamente",
            },
            "pendencias": ["Confirmar alíquota e cálculo do ICMS"],
            "estatisticas": {
                "tarifas_zona": len(tarifas),
                "faixas_peso": 5,
                "ufs": sorted({item["uf"] for item in tarifas.values()}),
                "zonas": sorted({item["zona"] for item in tarifas.values()}),
                "localidades": sum(len(itens) for itens in mapeamento.values()),
                "localidades_sem_cep": localidades_sem_cep,
            },
        }
    finally:
        workbook.close()
