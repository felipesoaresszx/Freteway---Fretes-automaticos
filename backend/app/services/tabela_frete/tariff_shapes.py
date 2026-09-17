"""Parsers plugaveis para formas semanticas de tabelas tarifarias."""
from __future__ import annotations

import json
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Protocol


def _key(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    return " ".join("".join(c for c in text if not unicodedata.combining(c)).upper().split())


def _number(value: object) -> float | None:
    if isinstance(value, (int, float)): return float(value)
    text = str(value or "").strip().replace("R$", "").replace(" ", "")
    if not text or not re.fullmatch(r"[\d.,]+", text): return None
    if "," in text: text = text.replace(".", "").replace(",", ".")
    return float(text)


def _days(value: object) -> int | None:
    values = [int(item) for item in re.findall(r"\d+", str(value or ""))]
    return max(values) if values else None


def _workbook_rows(path: Path) -> list[tuple[str, list[list[object]]]]:
    if path.suffix.lower() == ".xls":
        import xlrd
        workbook = xlrd.open_workbook(str(path))
        return [(sheet.name, [sheet.row_values(i) for i in range(sheet.nrows)]) for sheet in workbook.sheets()]
    from openpyxl import load_workbook
    workbook = load_workbook(path, data_only=True, read_only=True)
    try: return [(s.title, [list(r) for r in s.iter_rows(values_only=True)]) for s in workbook.worksheets]
    finally: workbook.close()


@dataclass(frozen=True)
class ShapeMatch:
    parser: str
    confidence: float
    data: dict
    issues: tuple[str, ...] = ()


class TariffShapeParser(Protocol):
    code: str
    def parse(self, path: Path, *, carrier: str | None = None) -> ShapeMatch | None: ...


REGION_LEVEL_ALIASES = {"POLO": "POLE", "CAPITAL": "POLE", "SEDE": "POLE", "METROPOLITANA": "POLE", "INTERIOR": "INTERIOR", "DEMAIS LOCALIDADES": "INTERIOR"}
STATE_NAMES = {"ACRE":"AC","ALAGOAS":"AL","AMAPA":"AP","AMAZONAS":"AM","BAHIA":"BA","CEARA":"CE","DISTRITO FEDERAL":"DF","ESPIRITO SANTO":"ES","GOIAS":"GO","MARANHAO":"MA","MATO GROSSO":"MT","MATO GROSSO DO SUL":"MS","MINAS GERAIS":"MG","PARA":"PA","PARAIBA":"PB","PARANA":"PR","PERNAMBUCO":"PE","PIAUI":"PI","RIO DE JANEIRO":"RJ","RIO GRANDE DO NORTE":"RN","RIO GRANDE DO SUL":"RS","RONDONIA":"RO","RORAIMA":"RR","SANTA CATARINA":"SC","SAO PAULO":"SP","SERGIPE":"SE","TOCANTINS":"TO"}


class PlaceCodeLegendParser:
    code = "place_code_region_legend"

    def parse(self, path: Path, *, carrier: str | None = None) -> ShapeMatch | None:
        if path.suffix.lower() not in {".xls", ".xlsx", ".xlsm"}: return None
        price_rows, legend = [], {}
        headers_matched = numeric_cells = numeric_valid = region_valid = 0
        for sheet_name, rows in _workbook_rows(path):
            for index, row in enumerate(rows):
                keys = [_key(cell) for cell in row]
                dcol = next((i for i,c in enumerate(keys) if c in {"DESTINO","PRACA","CODIGO","SIGLA"}), None)
                rcol = next((i for i,c in enumerate(keys) if c in {"REGIAO","NIVEL DE ATENDIMENTO"}), None)
                ecol = next((i for i,c in enumerate(keys) if "EXCEDENTE" in c and ("KG" in c or c == "EXCEDENTE")), None)
                bcol = next((i for i,c in enumerate(keys) if re.search(r"FRETE.*ATE\s*\d+\s*KG", c)), None)
                pcol = next((i for i,c in enumerate(keys) if "PRAZO" in c), None)
                if None not in {dcol, rcol, ecol, bcol}:
                    headers_matched = 4 + int(pcol is not None)
                    m = re.search(r"ATE\s*(\d+)\s*KG", keys[bcol]); limit = float(m.group(1)) if m else 100.0
                    for source_row, values in enumerate(rows[index+1:], start=index+2):
                        code = _key(values[dcol] if dcol < len(values) else ""); region = _key(values[rcol] if rcol < len(values) else "")
                        if not code or region not in REGION_LEVEL_ALIASES: break
                        base = _number(values[bcol] if bcol < len(values) else None); excess = _number(values[ecol] if ecol < len(values) else None)
                        numeric_cells += 2; numeric_valid += int(base is not None) + int(excess is not None); region_valid += 1
                        price_rows.append({"code":code,"region":region,"service_level":REGION_LEVEL_ALIASES[region],"base_weight_kg":limit,"base_price":base,"excess_rate":excess,"delivery_days":_days(values[pcol]) if pcol is not None and pcol < len(values) else None,"source":{"sheet":sheet_name,"row":source_row}})
                for col, cell in enumerate(keys):
                    if any(token in cell for token in ("SIGLAS","LEGENDA","UNIDADES")):
                        for values in rows[index+1:]:
                            code = _key(values[col] if col < len(values) else ""); label = _key(values[col+1] if col+1 < len(values) else "")
                            if not code or not label:
                                if legend: break
                                continue
                            if re.fullmatch(r"[A-Z0-9-]{2,10}", code): legend[code] = label
        if not price_rows: return None
        codes = {r["code"] for r in price_rows}; missing = sorted(codes - legend.keys()); index = _locality_index(path); destinations = []
        for row in price_rows:
            label = legend.get(row["code"]); city, state = _resolve_locality(label, index)
            # A linha INTERIOR usa a unidade apenas como polo operacional; ela
            # não representa a cidade da legenda como destino tarifário.
            if row["service_level"] == "INTERIOR": city = None
            destinations.append({"destination_code":row["code"],"legend_label":label,"uf":state,"city":city,"city_group":None,"cep_start":None,"cep_end":None,"region_code":row["region"],"service_level":row["service_level"],"delivery_days":row["delivery_days"],"weight_rates":[{"max_weight":row["base_weight_kg"],"price":row["base_price"],"raw":{}}] if row["base_price"] is not None else [],"tariff_rule":{"type":"BASE_PLUS_EXCESS","base_weight_kg":row["base_weight_kg"],"base_price":row["base_price"],"excess_rate_per_kg":row["excess_rate"]},"excess_weight_rate":row["excess_rate"],"fixed_surcharges":[],"percentage_surcharges":[],"cities":[city] if city else [],"source":row["source"]})
        confidence = .45 + .05*headers_matched + .15*(1-len(missing)/max(1,len(codes))) + .08*(region_valid/len(price_rows)) + .07*(numeric_valid/max(1,numeric_cells))
        resolved = sum(d["legend_label"] is not None and d["uf"] is not None for d in destinations)
        data = {"formato":"tabela_frete_universal_v1","shape":self.code,"carrier":carrier,"currency":"BRL","weight_bands":[],"destinations":destinations,"surcharges":[],"delivery_rules":[],"collection_rules":[],"general_rules":[],"destination_legend":{c:{"label":l,"scope":"TABLE"} for c,l in legend.items()},"region_level_aliases":REGION_LEVEL_ALIASES,"source_document":path.name,"estatisticas":{"pracas":len(destinations),"codigos":len(codes),"codigos_resolvidos":resolved}}
        return ShapeMatch(self.code, min(confidence,.99), data, ("destination_code_legend",) if missing or not legend else ())


def _locality_index(path: Path) -> dict[str, set[tuple[str,str]]]:
    candidates = [parent / "data" / "tariffs" / "rispa" / "resolucao_destino_cidades.json" for parent in path.parents]
    source = next((p for p in candidates if p.is_file()), None); result: dict[str,set[tuple[str,str]]] = {}
    if source:
        for item in json.loads(source.read_text(encoding="utf-8")):
            city, state = _key(item.get("cidade")), _key(item.get("uf"))
            if city and state: result.setdefault(city,set()).add((city,state))
    return result


def _resolve_locality(label: str | None, index: dict[str,set[tuple[str,str]]]) -> tuple[str|None,str|None]:
    normalized = _key(label)
    if normalized in STATE_NAMES: return None, STATE_NAMES[normalized]
    matches = index.get(normalized,set())
    if len(matches) == 1: return next(iter(matches))
    match = re.match(r"(.+?)[ /-]+([A-Z]{2})$", normalized)
    return (match.group(1),match.group(2)) if match else (normalized or None,None)


DEFAULT_SHAPE_PARSERS: tuple[TariffShapeParser,...] = (PlaceCodeLegendParser(),)


def parse_best_shape(path: Path, *, carrier: str | None = None) -> ShapeMatch | None:
    matches = [m for parser in DEFAULT_SHAPE_PARSERS if (m := parser.parse(path, carrier=carrier))]
    return max(matches, key=lambda item:item.confidence, default=None)
