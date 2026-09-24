"""Leitor semântico da tabela comercial Patrus para o contrato universal."""

from __future__ import annotations

import hashlib
import re
import unicodedata
from pathlib import Path


FORMAT = "tabela_frete_universal_v1"


def _key(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    return " ".join("".join(c for c in text if not unicodedata.combining(c)).upper().split())


def _cep(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    digits = re.sub(r"\D", "", str(value))
    return digits.zfill(8) if 1 <= len(digits) <= 8 else None


def _source(path: Path, sheet: str, row: int, field: str) -> dict:
    return {"source_document": path.name, "sheet": sheet, "row": row, "field": field}


def _find_sheet(workbook, *tokens: str):
    wanted = [_key(token) for token in tokens]
    return next((sheet for sheet in workbook.worksheets if all(token in _key(sheet.title) for token in wanted)), None)


def _find_header(sheet, required: set[str]) -> tuple[int, dict[str, int]]:
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
        headers = {_key(value): index for index, value in enumerate(row) if value is not None}
        if required.issubset(headers):
            return row_number, headers
    raise ValueError(f"Cabeçalho não encontrado na aba {sheet.title}")


def _locality_rows(sheet, *, start_col: int = 0, code: str | None = None) -> list[dict]:
    result = []
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
        cells = list(row[start_col:start_col + 6])
        normalized = {_key(value): index for index, value in enumerate(cells) if value is not None}
        if not {"UF", "LOCALIDADE", "CEP INICIAL", "CEP FINAL"}.issubset(normalized):
            continue
        indexes = normalized
        for data_row_number, values in enumerate(sheet.iter_rows(min_row=row_number + 1, values_only=True), row_number + 1):
            values = list(values[start_col:start_col + 6])
            uf = values[indexes["UF"]] if indexes["UF"] < len(values) else None
            city = values[indexes["LOCALIDADE"]] if indexes["LOCALIDADE"] < len(values) else None
            start = _cep(values[indexes["CEP INICIAL"]] if indexes["CEP INICIAL"] < len(values) else None)
            end = _cep(values[indexes["CEP FINAL"]] if indexes["CEP FINAL"] < len(values) else None)
            if not (uf and city and start and end):
                if result:
                    break
                continue
            result.append({
                "uf": str(uf).strip().upper(), "city": str(city).strip(),
                "cep_start": start, "cep_end": end, "code": code,
                "source": {"sheet": sheet.title, "row": data_row_number},
            })
        break
    return result


def _third_band_rows(sheet) -> list[dict]:
    groups = ((1, "MG - V. do Jequitinhonha", "MG"), (6, "BA - Barreiras (Oeste)", "BA"),
              (11, "SP - Litoral / V do Ribeira", "SP"), (16, "PE - Petrolina e Região", "PE"))
    result = []
    rows = list(sheet.iter_rows(values_only=True))
    header_index = next(i for i, row in enumerate(rows) if sum(_key(v) == "CEP INICIAL" for v in row) >= 3)
    for start_col, region, uf in groups:
        for row_number, row in enumerate(rows[header_index + 1:], header_index + 2):
            city = row[start_col + 1] if len(row) > start_col + 3 else None
            start = _cep(row[start_col + 2] if len(row) > start_col + 2 else None)
            end = _cep(row[start_col + 3] if len(row) > start_col + 3 else None)
            if city and start and end:
                result.append({"uf": uf, "city": str(city).strip(), "cep_start": start, "cep_end": end,
                               "region_code": region, "source": {"sheet": sheet.title, "row": row_number}})
    return result


def _cnpj_roots(sheet) -> list[dict]:
    header_row, headers = _find_header(sheet, {"RAIZ CNPJ", "RAZAO SOCIAL"})
    roots = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
        raw = row[headers["RAIZ CNPJ"]] if headers["RAIZ CNPJ"] < len(row) else None
        name = row[headers["RAZAO SOCIAL"]] if headers["RAZAO SOCIAL"] < len(row) else None
        digits = re.sub(r"\D", "", str(raw or ""))[:8]
        if len(digits) == 8 and name:
            roots.append({"root": digits, "name": str(name).strip(), "source": {"sheet": sheet.title, "row": row_number}})
    return roots


def is_patrus_workbook(path: str | Path) -> bool:
    if Path(path).suffix.lower() not in {".xlsx", ".xlsm"}:
        return False
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        names = {_key(sheet.title) for sheet in workbook.worksheets}
        return "TABELA FRACIONADA" in names and any("EMEX E TRT" in name for name in names) and "TAG" in names
    finally:
        workbook.close()


def extract_patrus_excel(path: str | Path) -> dict:
    from openpyxl import load_workbook

    file_path = Path(path)
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        tariff_sheet = _find_sheet(workbook, "Tabela Fracionada")
        general_sheet = _find_sheet(workbook, "Generalidades")
        capital_sheet = _find_sheet(workbook, "Capitais")
        third_sheet = _find_sheet(workbook, "Faixa", "Atend PE")
        emex_trt_sheet = _find_sheet(workbook, "EMEX", "TRT")
        tda_sheet = _find_sheet(workbook, "TDA")
        tag_sheet = _find_sheet(workbook, "TAG")
        no_group_sheet = _find_sheet(workbook, "Agrupamento")
        if not all((tariff_sheet, general_sheet, capital_sheet, third_sheet, emex_trt_sheet, tda_sheet, tag_sheet, no_group_sheet)):
            raise ValueError("A tabela Patrus não contém todas as abas obrigatórias")

        rows = list(tariff_sheet.iter_rows(values_only=True))
        header_index = next(i for i, row in enumerate(rows) if any(_key(v) == "DESTINO" for v in row) and any(_key(v) == "KG EXCED." for v in row))
        header = rows[header_index]
        destination_col = next(i for i, v in enumerate(header) if _key(v) == "DESTINO")
        excess_col = next(i for i, v in enumerate(header) if _key(v) == "KG EXCED.")
        tso_col = next(i for i, v in enumerate(header) if "SEGURO OBRIGATORIO" in _key(v))
        fvm_col = next(i for i, v in enumerate(header) if _key(v).startswith("FRETE VALOR"))
        limits = [(i, float(v)) for i, v in enumerate(rows[header_index + 1]) if destination_col < i < excess_col and isinstance(v, (int, float)) and v > 0]
        destinations, by_region = [], {}
        for row_number, row in enumerate(rows[header_index + 3:], header_index + 4):
            label = row[destination_col] if destination_col < len(row) else None
            match = re.match(r"\s*([A-Z]{2})\s*-\s*(.+)", str(label or ""))
            if not match:
                continue
            uf, kind = match.groups()
            unavailable_text = " ".join(str(row[col] or "") for col, _ in limits)
            status = "SOB_CONSULTA" if "SOB CONSULTA" in _key(unavailable_text) else "INDISPONIVEL_POR_TABELA" if "SEM NEG" in _key(unavailable_text) else "AVAILABLE"
            bands = []
            lower = 0.0
            for col, upper in limits:
                value = row[col] if col < len(row) else None
                if isinstance(value, (int, float)) and value > 0:
                    bands.append({"min_weight": lower, "max_weight": upper, "price": float(value),
                                  "raw": _source(file_path, tariff_sheet.title, row_number, f"peso_ate_{upper:g}")})
                lower = upper
            region_code = f"{uf} - {kind.strip()}"
            service_level = "INTERIOR" if "INTERIOR" in _key(kind) else "CAPITAL" if "CAPITAL" in _key(kind) else "SPECIAL"
            item = {"uf": uf, "city": None, "region_code": region_code, "type": kind.strip().upper(), "service_level": service_level,
                    "status": status, "status_detail": unavailable_text.strip() or None, "weight_rates": bands,
                    "excess_weight_rate": float(row[excess_col]) if excess_col < len(row) and isinstance(row[excess_col], (int, float)) and row[excess_col] > 0 else None,
                    "regional_surcharges": [
                        {"code": "TSO", "type": "PERCENTAGE", "basis": "INVOICE_VALUE", "value": float(row[tso_col]), "minimum": 9.0}
                        for _ in [0] if tso_col < len(row) and isinstance(row[tso_col], (int, float)) and row[tso_col] > 0
                    ] + [
                        {"code": "FVM", "type": "PERCENTAGE", "basis": "INVOICE_VALUE", "value": float(row[fvm_col]), "minimum": 8.36}
                        for _ in [0] if fvm_col < len(row) and isinstance(row[fvm_col], (int, float)) and row[fvm_col] > 0
                    ], "source": _source(file_path, tariff_sheet.title, row_number, "tarifa_regiao")}
            destinations.append(item)
            by_region[_key(region_code)] = item

        # Algumas descrições de indisponibilidade aparecem em célula mesclada
        # somente na primeira linha do estado (por exemplo, Ceará).
        for index, item in enumerate(destinations[1:], 1):
            previous = destinations[index - 1]
            if item["uf"] == previous["uf"] and not item["weight_rates"] and item["status"] == "AVAILABLE" and previous["status"] != "AVAILABLE":
                item["status"] = previous["status"]
                item["status_detail"] = previous["status_detail"]

        def specialized(record: dict) -> dict:
            template = by_region.get(_key(record["region_code"]))
            if not template:
                raise ValueError(f"Região sem tarifa: {record['region_code']}")
            return {**template, **record, "weight_rates": list(template["weight_rates"]),
                    "regional_surcharges": list(template["regional_surcharges"])}

        capital_rows = _locality_rows(capital_sheet)
        for record in capital_rows:
            region = next((r for r in destinations if r["uf"] == record["uf"] and "CAPITAL" in r["type"]), None)
            if region:
                destinations.append({**region, **record, "region_code": region["region_code"],
                                     "weight_rates": list(region["weight_rates"]), "regional_surcharges": list(region["regional_surcharges"])})
        third_rows = _third_band_rows(third_sheet)
        destinations.extend(specialized(record) for record in third_rows)

        # Listas de incidência permanecem dados, sem cidades hardcoded.
        emex = _locality_rows(emex_trt_sheet, start_col=0, code="EMEX")
        trt = _locality_rows(emex_trt_sheet, start_col=6, code="TRT")
        tda = _locality_rows(tda_sheet, code="TDA")
        tag = _cnpj_roots(tag_sheet)
        no_group = _cnpj_roots(no_group_sheet)
        rules = [
            {"code": "CTE", "name": "Taxa por CT-e", "type": "FIXED", "value": 5.02},
            {"code": "PEDAGIO", "name": "Pedágio", "type": "WEIGHT_FRACTION", "value": 8.22, "fraction_kg": 100},
            {"code": "GRIS", "name": "Gerenciamento de risco", "type": "PERCENTAGE", "basis": "INVOICE_VALUE", "value": .0017, "minimum": 7.45},
            {"code": "TRT", "name": "Restrição de trânsito", "type": "PERCENTAGE", "basis": "ORIGINAL_FREIGHT", "value": .1109, "minimum": 22.18, "cep_ranges": trt},
            {"code": "EMEX", "name": "Emergência excepcional", "type": "PERCENTAGE", "basis": "INVOICE_VALUE", "value": .0033, "minimum": 34.74, "cep_ranges": emex},
            {"code": "TDA", "name": "Dificuldade de acesso", "type": "PERCENTAGE", "basis": "ORIGINAL_FREIGHT", "value": .1663, "minimum": 47.93, "cep_ranges": tda},
            {"code": "TDE", "status": "UNRESOLVED", "value": .1109, "minimum": 182.36, "reason": "Lista de incidência ausente no arquivo"},
            {"code": "TDE2", "status": "UNRESOLVED", "value": .3326, "minimum": 580.17, "reason": "Lista de incidência ausente no arquivo"},
            {"code": "TDE3", "status": "UNRESOLVED", "value": .4435, "minimum": 1047.68, "reason": "Lista de incidência ausente no arquivo"},
            {"code": "TAG", "name": "Taxa de agendamento", "status": "OPERATIONAL",
             "type": "PERCENTAGE", "basis": "ORIGINAL_FREIGHT", "value": 0.0, "minimum": 0.0,
             "cnpj_roots": tag, "reason": "Tabela comercial informa alíquota e mínimo iguais a zero",
             "source": _source(file_path, general_sheet.title, 38, "TAG")},
            {"code": "NAO_AGRUPAMENTO", "status": "OPERATIONAL", "cnpj_roots": no_group},
            {"code": "REENTREGA", "status": "OPTIONAL", "type": "PERCENTAGE", "basis": "ORIGINAL_FREIGHT", "value": .5, "minimum": 62.10},
            {"code": "CAP", "status": "OPTIONAL", "minimum": 96.84, "reason": "Depende de solicitação de paletização"},
        ]
        version = next((str(v) for row in rows[:10] for v in row if isinstance(v, str) and re.search(r"V\s*\d", v, re.I)), None)
        digest = hashlib.sha256(file_path.read_bytes()).hexdigest()
        pending = [rule for rule in rules if rule.get("status") == "UNRESOLVED"]
        return {
            "formato": FORMAT, "carrier": "Patrus", "carrier_legal_name": "Patrus Transportes Ltda",
            "origin": {"city": "São Paulo", "state": "SP"}, "currency": "BRL", "fator_cubagem": 300,
            # O contrato orienta solicitar o prazo à unidade responsável e não
            # publica quantidade de dias. Zero mantém a cotação disponível no
            # contrato Sankhya sem inventar um prazo comercial.
            "default_delivery_days": 0,
            "weight_policy": "max_real_cubed", "destinations": destinations, "pracas": destinations,
            "surcharges": rules, "faixas_tarifarias": [band for item in destinations[:len(by_region)] for band in item["weight_rates"]],
            "tax_rules": [{
                "code": "ICMS", "name": "ICMS sobre transporte", "type": "GROSS_UP",
                "rates_by_destination": {
                    "BA": .07, "CE": .07, "ES": .07, "MG": .12, "PE": .07,
                    "PR": .12, "RJ": .12, "RS": .12, "SC": .12, "SE": .07, "SP": .12,
                },
                "source": _source(file_path, tariff_sheet.title, 67, "Impostos ICMS e/ou ISS oficiais, não inclusos"),
            }],
            "special_lists": {"capitals": capital_rows, "third_band": third_rows, "emex": emex, "trt": trt,
                              "tda": tda, "tag": tag, "non_grouping": no_group},
            "operational_rules": {"pallet_weight_kg": 1000, "redelivery": {"percentage": .5, "minimum": 62.10}},
            "unresolved_rules": pending, "table_version": version, "source_document": file_path.name, "source_sha256": digest,
            "validation": {"status": "NEEDS_REVIEW" if pending else "TABLE_VALIDATED",
                           "errors": [f"{r['code']}: {r['reason']}" for r in pending], "warnings": []},
            "estatisticas": {"regions": len(by_region), "destinations": len(destinations), "cep_ranges": len(destinations) - len(by_region),
                             "tda": len(tda), "trt": len(trt), "emex": len(emex), "tag": len(tag), "non_grouping": len(no_group)},
        }
    finally:
        workbook.close()
