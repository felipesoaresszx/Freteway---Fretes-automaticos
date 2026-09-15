from __future__ import annotations

from pathlib import Path


class DocumentExtractor:
    def extract(self, path: str | Path) -> str:
        p = Path(path)
        if p.suffix.lower() == ".csv":
            return p.read_text(encoding="utf-8-sig")
        if p.suffix.lower() in {".xlsx", ".xlsm", ".xls"}:
            from openpyxl import load_workbook

            workbook = load_workbook(p, data_only=True, read_only=True)
            chunks: list[str] = []
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    values = [str(v).strip() for v in row if v is not None and str(v).strip()]
                    if values:
                        chunks.append(" | ".join(values))
            workbook.close()
            return "\n".join(chunks)
        return p.read_text(encoding="utf-8-sig", errors="ignore")


def extract_document(path: str | Path) -> str:
    return DocumentExtractor().extract(path)
