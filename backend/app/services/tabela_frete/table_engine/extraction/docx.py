from pathlib import Path


def extract_docx(path: str | Path) -> str:
    from docx import Document

    document = Document(str(path))
    chunks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    chunks.extend(
        " | ".join(cell.text.strip() for cell in row.cells)
        for table in document.tables
        for row in table.rows
        if any(cell.text.strip() for cell in row.cells)
    )
    return "\n".join(chunks)
