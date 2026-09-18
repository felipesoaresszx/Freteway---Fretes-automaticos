from __future__ import annotations

from pathlib import Path


def extract_excel(path: str | Path) -> str:
    p = Path(path)
    chunks: list[str] = []
    if p.suffix.lower() == ".xls":
        import xlrd

        workbook = xlrd.open_workbook(str(p))
        for sheet in workbook.sheets():
            chunks.append(f"### {sheet.name}")
            for row_index in range(sheet.nrows):
                values = [str(v).strip() for v in sheet.row_values(row_index) if str(v).strip()]
                if values:
                    chunks.append(" | ".join(values))
    else:
        from openpyxl import load_workbook

        workbook = load_workbook(p, data_only=True, read_only=True)
        for sheet in workbook.worksheets:
            chunks.append(f"### {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = [str(v).strip() for v in row if v is not None and str(v).strip()]
                if values:
                    chunks.append(" | ".join(values))
        workbook.close()
    return "\n".join(chunks)
