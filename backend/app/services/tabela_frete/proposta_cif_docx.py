"""Parser de propostas CIF em DOCX com tarifa por kg, percentual e minimo."""

from __future__ import annotations

import re
import unicodedata
from pathlib import Path

from app.services.tabela_frete.table_engine.normalization.city_normalizer import (
    normalize_city_name,
)
from app.services.tabela_frete.table_engine.normalization.number_normalizer import (
    normalize_number,
)
from app.services.tabela_frete.regioes_imediatas_ibge import REGIOES_IMEDIATAS_IBGE


FORMAT = "tabela_frete_universal_v1"


def _key(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^A-Z0-9%]+", " ", text.upper()).strip()


def _money(value: object) -> float | None:
    return normalize_number(value)


def _percentage(value: object) -> float | None:
    number = normalize_number(value)
    if number is None:
        return None
    return number / 100 if number > 1 else number


def _delivery_range(value: object) -> tuple[int | None, int | None]:
    days = [int(item) for item in re.findall(r"\d+", str(value or ""))]
    if not days:
        return None, None
    return min(days), max(days)


def _destination(value: object) -> tuple[str, str] | None:
    text = " ".join(str(value or "").split())
    match = re.match(r"^(.+?)\s*[-\u2013\u2014\ufffd]\s*([A-Za-z]{2})\s*$", text)
    if not match:
        return None
    city = normalize_city_name(match.group(1))["normalized_value"]
    uf = match.group(2).upper()
    if not city or len(uf) != 2:
        return None
    return city, uf


def _column_indexes(headers: list[str]) -> dict[str, int] | None:
    indexes: dict[str, int] = {}
    for index, header in enumerate(headers):
        key = _key(header)
        if key == "DESTINO":
            indexes["destination"] = index
        elif "FRETE" in key and "PESO" in key:
            indexes["weight_rate"] = index
        elif "FRETE" in key and "%" in key:
            indexes["freight_percentage"] = index
        elif "FRETE" in key and "MINIMO" in key:
            indexes["minimum_freight"] = index
        elif "PRAZO" in key and "ENTREGA" in key:
            indexes["delivery_days"] = index
    required = {
        "destination", "weight_rate", "freight_percentage",
        "minimum_freight", "delivery_days",
    }
    return indexes if required <= indexes.keys() else None


def _cell(row: list[str], indexes: dict[str, int], name: str) -> str:
    index = indexes[name]
    return row[index].strip() if index < len(row) else ""


def _tariff(row: list[str], indexes: dict[str, int]) -> dict[str, float] | None:
    per_kg = _money(_cell(row, indexes, "weight_rate"))
    percentage = _percentage(_cell(row, indexes, "freight_percentage"))
    minimum = _money(_cell(row, indexes, "minimum_freight"))
    if per_kg is None or percentage is None or minimum is None:
        return None
    return {
        "type": "BASE_PLUS_EXCESS",
        "base_weight_kg": 0.0,
        "base_price": 0.0,
        "excess_rate_per_kg": per_kg,
        "minimum_freight": minimum,
        "freight_percentage": percentage,
    }


def parse_cif_proposal_rows(
    rows: list[list[str]], *, source_document: str, full_text: str = "",
) -> dict | None:
    """Converte a tabela da proposta sem ampliar a cobertura descrita no arquivo."""
    if len(rows) < 2:
        return None
    indexes = _column_indexes(rows[0])
    if indexes is None:
        return None

    destinations: list[dict] = []
    unresolved_regions: list[dict] = []
    mapped_regions: list[dict] = []
    last_named_destination: dict | None = None
    for position, row in enumerate(rows[1:], start=2):
        label = _cell(row, indexes, "destination")
        tariff = _tariff(row, indexes)
        minimum_days, maximum_days = _delivery_range(_cell(row, indexes, "delivery_days"))
        named = _destination(label)
        if named and tariff and maximum_days is not None:
            city, uf = named
            destination = {
                "origin_uf": "SP",
                "origin_city": "GUARULHOS",
                "uf": uf,
                "city": city,
                "cities": [city],
                "delivery_days": maximum_days,
                "delivery_days_min": minimum_days,
                "weight_rates": [],
                "tariff_rule": tariff,
                "minimum_freight": tariff["minimum_freight"],
                "freight_percentage": tariff["freight_percentage"],
                "cubage_factor": 300.0,
                "conditions": {"requires_city_match": True},
                "source": {"document": source_document, "table_row": position},
            }
            destinations.append(destination)
            last_named_destination = destination
            continue

        if _key(label) == "REGIAO":
            anchor_city = last_named_destination.get("city") if last_named_destination else None
            uf = last_named_destination.get("uf") if last_named_destination else None
            region = REGIOES_IMEDIATAS_IBGE.get((uf, anchor_city))
            region_data = {
                "label": label,
                "reference_city": anchor_city,
                "uf": uf,
                "delivery_days_min": minimum_days,
                "delivery_days_max": maximum_days,
                "tariff_rule": tariff,
                "source": {"document": source_document, "table_row": position},
            }
            if region and tariff and maximum_days is not None:
                cities = [normalize_city_name(city)["normalized_value"] for city in region["municipios"]]
                region_code = f"IBGE_IMEDIATA_{region['id']}"
                destinations.append({
                    "origin_uf": "SP",
                    "origin_city": "GUARULHOS",
                    "uf": uf,
                    "city": None,
                    "cities": cities,
                    "cep_start": region["cep_start"],
                    "cep_end": region["cep_end"],
                    "region_code": region_code,
                    "service_level": "REGIAO_IMEDIATA",
                    "delivery_days": maximum_days,
                    "delivery_days_min": minimum_days,
                    "weight_rates": [],
                    "tariff_rule": tariff,
                    "minimum_freight": tariff["minimum_freight"],
                    "freight_percentage": tariff["freight_percentage"],
                    "cubage_factor": 300.0,
                    "conditions": {"requires_city_match": True},
                    "source": {
                        "document": source_document,
                        "table_row": position,
                        "coverage": "IBGE_REGIAO_GEOGRAFICA_IMEDIATA",
                        "ibge_region_id": region["id"],
                        "commercial_cep_range": True,
                    },
                })
                mapped_regions.append({
                    **region_data,
                    "region_code": region_code,
                    "ibge_region_id": region["id"],
                    "cep_start": region["cep_start"],
                    "cep_end": region["cep_end"],
                    "cities": cities,
                    "status": "MAPPED",
                })
            else:
                unresolved_regions.append({**region_data, "status": "UNRESOLVED_COVERAGE"})

    if not destinations:
        return None

    normalized_text = _key(full_text)
    insurance_match = re.search(r"\bSEGURO\s+(\d+(?:[,.]\d+)?)\s*%", full_text, re.I)
    surcharges = []
    if insurance_match:
        rate = float(insurance_match.group(1).replace(",", ".")) / 100
        surcharges.append({
            "code": "SEGURO",
            "name": "Seguro",
            "type": "PERCENTAGE",
            "value": rate,
            "minimum": None,
            "basis": "INVOICE_VALUE",
            "conditions": {},
            "source": source_document,
        })

    general_rules = [{
        "kind": "commercial_proposal",
        "cubage_factor": 300 if re.search(r"CUBAGEM\s*300", normalized_text) else None,
        "return_percentage": 1.0 if "DEVOLUCAO 100%" in normalized_text else None,
        "redelivery_percentage": .5 if "REENTREGA 50%" in normalized_text else None,
        "icms": "RESOLVED_DEFAULT_7_PERCENT",
        "source": source_document,
    }]
    tariff_rules = [
        dict(
            item["tariff_rule"], city=item.get("city"), uf=item["uf"],
            region_code=item.get("region_code"),
        )
        for item in destinations
    ]
    return {
        "formato": FORMAT,
        "canonical_schema": "canonical_tariff_v2",
        "schema_version": 2,
        "origin": {"city": "GUARULHOS", "state": "SP"},
        "currency": "BRL",
        "fator_cubagem": 300.0,
        "weight_bands": [],
        "destinations": destinations,
        "pracas": destinations,
        "faixas_tarifarias": tariff_rules,
        "surcharges": surcharges,
        "tax_rules": [{
            "code": "ICMS",
            "name": "ICMS padrão",
            "type": "GROSS_UP",
            "rates_by_destination": {"*": .07},
            "default_rate": .07,
            "source": source_document,
        }],
        "delivery_rules": [],
        "collection_rules": [],
        "general_rules": general_rules,
        "mapped_regions": mapped_regions,
        "unresolved_regions": unresolved_regions,
        "metadata": {"source_document": source_document, "parser": "proposta_cif_docx_v1"},
        "source_document": source_document,
        "estatisticas": {
            "pracas": len(destinations),
            "regras_tarifarias": len(tariff_rules),
            "regioes_mapeadas": len(mapped_regions),
            "regioes_sem_malha": len(unresolved_regions),
        },
    }


def extract_cif_proposal_docx(path: str | Path) -> dict | None:
    from docx import Document

    document = Document(str(path))
    full_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    for table in document.tables:
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        parsed = parse_cif_proposal_rows(
            rows, source_document=Path(path).name, full_text=full_text,
        )
        if parsed:
            return parsed
    return None
