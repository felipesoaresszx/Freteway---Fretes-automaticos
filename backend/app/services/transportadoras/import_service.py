import csv, io, logging, re, time, unicodedata
from collections import Counter
from datetime import datetime
from pathlib import Path
from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook
from sqlalchemy import or_, select
from sqlalchemy.ext.asyncio import AsyncSession
from app.core.config import get_settings
from app.models.models import Transportadora, TransportadoraFonte, TransportadoraImportacao, TransportadoraImportacaoItem
from .normalization import normalize_bool, normalize_cep, normalize_cnpj, normalize_method, normalize_phone, normalize_status, normalize_text, normalize_uf, normalize_url, validate_cnpj

logger=logging.getLogger(__name__)
FIELDS={"codigo_importacao","nome_fantasia","razao_social","cnpj","situacao_rntrc","categoria_rntrc","status_cnpj","metodo_atual","integracao_disponivel","qtde_referencia","site","portal_cliente_cotacao","api_documentacao","email_comercial","telefone","logradouro","numero","complemento","bairro","cidade","uf","cep","cnae_principal","rntrc","cobertura_resumo","status_validacao","precisa_revisao","observacoes","fonte_principal","ativa"}
ALIASES={**{field:field for field in FIELDS},"nome":"nome_fantasia","fantasia":"nome_fantasia","transportadora":"nome_fantasia","nome_empresarial":"razao_social","nome_transportador":"razao_social","cnpj_transportadora":"cnpj","cpfcnpjtransportador":"cnpj","cpf_cnpj_transportador":"cnpj","numero_rntrc":"rntrc","situacao":"situacao_rntrc","categoria_transportador":"categoria_rntrc","categoria":"categoria_rntrc","municipio":"cidade","estado":"uf","e_mail":"email_comercial","email":"email_comercial","portal_cotacao":"portal_cliente_cotacao"}
FIELD_MAP={"codigo_importacao":"codigo_importacao","nome_fantasia":"nome","razao_social":"razao_social","cnpj":"cnpj_cpf","status_cnpj":"status_cnpj","integracao_disponivel":"integracao_disponivel","site":"site","portal_cliente_cotacao":"portal_cotacao","api_documentacao":"api_documentacao","email_comercial":"email_comercial","telefone":"telefone","logradouro":"logradouro","numero":"numero","complemento":"complemento","bairro":"bairro","cidade":"cidade","uf":"uf","cep":"cep","cnae_principal":"cnae_principal","rntrc":"rntrc","cobertura_resumo":"cobertura_resumo","precisa_revisao":"precisa_revisao","status_validacao":"status_validacao","observacoes":"observacoes"}

def canonical_header(value:object)->str:
    text=normalize_text(value) or ""; text="".join(c for c in unicodedata.normalize("NFD",text) if unicodedata.category(c)!="Mn")
    key=re.sub(r"[^a-z0-9]+","_",text.lower()).strip("_"); return ALIASES.get(key,key)

def _dict_rows(headers:list[str],rows:list[list])->list[dict]:
    return [{headers[i]:v for i,v in enumerate(row) if i<len(headers) and headers[i]} for row in rows if any(v not in (None,"") for v in row)]

def _rows_from_xlsx(content:bytes)->tuple[list[dict],list[dict]]:
    try: workbook=load_workbook(io.BytesIO(content),read_only=True,data_only=True,keep_links=False)
    except Exception as exc: raise HTTPException(422,"Arquivo XLSX inválido ou corrompido") from exc
    main="IMPORT_FRETEWAY" if "IMPORT_FRETEWAY" in workbook.sheetnames else workbook.sheetnames[0]
    def read(name):
        values=list(workbook[name].iter_rows(values_only=True)); return _dict_rows([canonical_header(v) for v in values[0]],[list(row) for row in values[1:]]) if values else []
    return read(main),read("FONTES_E_VALIDACAO") if "FONTES_E_VALIDACAO" in workbook.sheetnames else []

def _rows_from_xls(content:bytes)->tuple[list[dict],list[dict]]:
    try:
        import xlrd
        workbook=xlrd.open_workbook(file_contents=content,on_demand=True)
    except ModuleNotFoundError as exc: raise HTTPException(503,"Leitura de XLS indisponível; instale as dependências do backend") from exc
    except Exception as exc: raise HTTPException(422,"Arquivo XLS inválido ou corrompido") from exc
    name="IMPORT_FRETEWAY" if "IMPORT_FRETEWAY" in workbook.sheet_names() else workbook.sheet_names()[0]; sheet=workbook.sheet_by_name(name)
    return (_dict_rows([canonical_header(v) for v in sheet.row_values(0)],[sheet.row_values(i) for i in range(1,sheet.nrows)]),[]) if sheet.nrows else ([],[])

def _rows_from_csv(content:bytes)->tuple[list[dict],list[dict]]:
    try: text=content.decode("utf-8-sig")
    except UnicodeDecodeError as exc: raise HTTPException(422,"CSV deve estar em UTF-8 ou UTF-8-SIG") from exc
    try:
        dialect=csv.Sniffer().sniff(text[:4096],delimiters=",;"); return [{canonical_header(k):v for k,v in row.items() if k} for row in csv.DictReader(io.StringIO(text),dialect=dialect)],[]
    except csv.Error as exc: raise HTTPException(422,"CSV inválido") from exc

async def parse_upload(file:UploadFile)->tuple[list[dict],list[dict]]:
    suffix=Path(file.filename or "").suffix.lower()
    if suffix not in {".xlsx",".xls",".csv"}: raise HTTPException(415,"Envie um arquivo .xlsx, .xls ou .csv")
    max_size=get_settings().TRANSPORTADORA_IMPORT_MAX_BYTES; content=await file.read(max_size+1)
    if not content: raise HTTPException(422,"Arquivo vazio")
    if len(content)>max_size: raise HTTPException(413,f"Arquivo excede {max_size//(1024*1024)} MB")
    mime=(file.content_type or "").split(";",1)[0].lower(); allowed={".csv":{"text/csv","text/plain","application/csv","application/vnd.ms-excel","application/octet-stream"},".xlsx":{"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet","application/zip","application/octet-stream"},".xls":{"application/vnd.ms-excel","application/octet-stream"}}
    if mime and mime not in allowed[suffix]: raise HTTPException(415,"O conteúdo do arquivo não corresponde ao formato informado")
    if suffix==".xlsx" and not content.startswith(b"PK"): raise HTTPException(422,"Arquivo XLSX inválido")
    if suffix==".xls" and not content.startswith(bytes.fromhex("D0CF11E0")): raise HTTPException(422,"Arquivo XLS inválido")
    rows,sources=_rows_from_xlsx(content) if suffix==".xlsx" else _rows_from_xls(content) if suffix==".xls" else _rows_from_csv(content)
    if not rows: raise HTTPException(422,"Arquivo não possui registros")
    if not ({"cnpj","nome_fantasia","razao_social"}&set().union(*(row.keys() for row in rows[:20]))): raise HTTPException(422,"Informe CNPJ ou nome/razão social")
    return rows,sources

def normalize_row(row:dict)->tuple[dict,list[str],list[str]]:
    data={key:normalize_text(row.get(key)) for key in FIELDS}; data.update(cnpj=normalize_cnpj(row.get("cnpj")),cep=normalize_cep(row.get("cep")),telefone=normalize_phone(row.get("telefone")),uf=normalize_uf(row.get("uf")),site=normalize_url(row.get("site")),portal_cliente_cotacao=normalize_url(row.get("portal_cliente_cotacao")),api_documentacao=normalize_url(row.get("api_documentacao")),fonte_principal=normalize_url(row.get("fonte_principal")),precisa_revisao=normalize_bool(row.get("precisa_revisao")),status_validacao=normalize_status(row.get("status_validacao")),metodo_atual=normalize_method(row.get("metodo_atual")))
    data["nome_fantasia"]=data["nome_fantasia"] or data["razao_social"]; data["razao_social"]=data["razao_social"] or data["nome_fantasia"]; data["rntrc"]=re.sub(r"\D","",data["rntrc"] or "") or None
    errors=[]; warnings=[]
    if not data["cnpj"] and not (data["nome_fantasia"] and data["cidade"] and data["uf"]): errors.append("Informe CNPJ ou nome com cidade e UF")
    if data["cnpj"] and not validate_cnpj(data["cnpj"]): errors.append("CNPJ inválido")
    if data["cep"] and len(data["cep"])!=8: errors.append("CEP deve possuir 8 dígitos")
    if row.get("uf") and not data["uf"]: errors.append("UF inválida")
    for field in ("site","portal_cliente_cotacao","api_documentacao","fonte_principal"):
        if row.get(field) and not data[field]: warnings.append(f"URL ignorada por formato inválido: {field}")
    if data["email_comercial"] and not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+",data["email_comercial"]): data["email_comercial"]=None; warnings.append("E-mail ignorado por formato inválido")
    if data["telefone"] and len(data["telefone"]) not in range(10,14): data["telefone"]=None; warnings.append("Telefone ignorado por formato inválido")
    if not data["cnpj"]: warnings.append("CNPJ ausente; registro exige revisão manual")
    return data,errors,warnings

async def create_preview(db:AsyncSession,file:UploadFile,user_id:str|None)->TransportadoraImportacao:
    started=time.monotonic(); rows,sources=await parse_upload(file); normalized=[normalize_row(row) for row in rows]
    cnpjs=[d["cnpj"] for d,_,_ in normalized if d["cnpj"]]; rntrcs=[d["rntrc"] for d,_,_ in normalized if d["rntrc"]]; names=[d["nome_fantasia"] for d,_,_ in normalized if d["nome_fantasia"]]
    carriers=list((await db.scalars(select(Transportadora).where(or_(Transportadora.cnpj_cpf.in_(cnpjs),Transportadora.rntrc.in_(rntrcs),Transportadora.nome.in_(names)),Transportadora.deleted_at.is_(None)))).all())
    by_cnpj={x.cnpj_cpf:x for x in carriers if x.cnpj_cpf}; by_rntrc={x.rntrc:x for x in carriers if x.rntrc}; by_identity={(x.nome.casefold(),(x.cidade or "").casefold(),x.uf or ""):x for x in carriers}; cnpj_counts=Counter(cnpjs); rntrc_counts=Counter(rntrcs)
    operation=TransportadoraImportacao(id="imp_"+datetime.utcnow().strftime("%Y%m%d%H%M%S%f"),arquivo_nome=Path(file.filename or "arquivo").name,usuario_id=user_id,total_registros=len(rows),origem="FILE",fontes=[{k:str(v) if v is not None else None for k,v in source.items()} for source in sources]); db.add(operation)
    for line,(original,(data,errors,warnings)) in enumerate(zip(rows,normalized),2):
        cnpj=data["cnpj"]; rntrc=data["rntrc"]; identity=((data["nome_fantasia"] or "").casefold(),(data["cidade"] or "").casefold(),data["uf"] or ""); carrier=by_cnpj.get(cnpj) or by_rntrc.get(rntrc) or by_identity.get(identity)
        if errors: outcome="ERRO"
        elif cnpj and cnpj_counts[cnpj]>1: outcome="DUPLICADO"; errors.append("CNPJ duplicado no arquivo")
        elif rntrc and rntrc_counts[rntrc]>1: outcome="DUPLICADO"; errors.append("RNTRC duplicado no arquivo")
        elif carrier: outcome="ATUALIZACAO" if any(v is not None and getattr(carrier,target,None)!=v for source,target in FIELD_MAP.items() if (v:=data.get(source)) is not None) else "IGNORADO"
        elif data["precisa_revisao"] or not cnpj: outcome="REVISAO"
        else: outcome="NOVO"
        db.add(TransportadoraImportacaoItem(importacao=operation,linha=line,codigo_importacao=data["codigo_importacao"],cnpj=cnpj,nome_transportadora=data["nome_fantasia"],resultado=outcome,dados_originais={k:str(v) if v is not None else None for k,v in original.items()},dados_normalizados=data,erros=errors,avisos=warnings))
    await db.flush(); totals=Counter(item.resultado for item in operation.itens); operation.novos=totals["NOVO"]; operation.atualizados=totals["ATUALIZACAO"]; operation.ignorados=totals["IGNORADO"]+totals["DUPLICADO"]; operation.erros=totals["ERRO"]; operation.revisao=totals["REVISAO"]; await db.commit(); logger.info("import_preview id=%s rows=%s duration_ms=%s",operation.id,len(rows),round((time.monotonic()-started)*1000)); return operation

async def confirm_import(db:AsyncSession,operation:TransportadoraImportacao,update_existing:bool,import_review:bool,activate_imported:bool=False)->tuple[dict,list[str]]:
    if operation.status!="PREVIEW": raise HTTPException(409,"Importação já confirmada ou indisponível")
    cnpjs=[i.cnpj for i in operation.itens if i.cnpj]; rntrcs=[i.dados_normalizados.get("rntrc") for i in operation.itens if i.dados_normalizados.get("rntrc")]; names=[i.nome_transportadora for i in operation.itens if i.nome_transportadora]
    carriers=list((await db.scalars(select(Transportadora).where(or_(Transportadora.cnpj_cpf.in_(cnpjs),Transportadora.rntrc.in_(rntrcs),Transportadora.nome.in_(names))))).all()); by_cnpj={x.cnpj_cpf:x for x in carriers if x.cnpj_cpf}; by_rntrc={x.rntrc:x for x in carriers if x.rntrc}; by_identity={(x.nome.casefold(),(x.cidade or "").casefold(),x.uf or ""):x for x in carriers}; result=Counter(); created_ids=[]
    try:
        operation.status="IMPORTING"
        for item in operation.itens:
            if item.resultado in {"ERRO","DUPLICADO"}: result["erros" if item.resultado=="ERRO" else "ignorados"]+=1; continue
            if item.resultado=="REVISAO" and not import_review: result["revisao"]+=1; continue
            data=item.dados_normalizados; identity=((item.nome_transportadora or "").casefold(),(data.get("cidade") or "").casefold(),data.get("uf") or ""); carrier=by_cnpj.get(item.cnpj) or by_rntrc.get(data.get("rntrc")) or by_identity.get(identity)
            if carrier and not update_existing: result["ignorados"]+=1; item.transportadora_id=carrier.id; continue
            if carrier:
                for source,target in FIELD_MAP.items():
                    if data.get(source) is not None: setattr(carrier,target,data[source])
                result["atualizados"]+=1
            else:
                values={target:data.get(source) for source,target in FIELD_MAP.items() if data.get(source) is not None}; values["razao_social"]=values.get("razao_social") or values.get("nome"); carrier=Transportadora(**values,segmento="Transportadora",tipo_integracao="n8n",metodo_calculo="manual",status_integracao="nao_aplicavel",ativa=activate_imported,taxa_sucesso=0,tempo_medio_ms=0,origem_cadastro="XLSX" if operation.arquivo_nome.lower().endswith((".xlsx",".xls")) else "CSV",imported_at=datetime.utcnow(),enrichment_status="NOT_STARTED"); db.add(carrier); await db.flush(); result["criados"]+=1; created_ids.append(carrier.id)
            item.transportadora_id=carrier.id
            if data.get("fonte_principal"): db.add(TransportadoraFonte(transportadora_id=carrier.id,tipo_fonte="IMPORTACAO",url=data["fonte_principal"],descricao="Fonte principal informada no arquivo"))
        operation.status="SUCCESS" if not result["erros"] else "PARTIAL"; operation.finished_at=datetime.utcnow(); operation.novos=result["criados"]; operation.atualizados=result["atualizados"]; operation.ignorados=result["ignorados"]; operation.erros=result["erros"]; operation.revisao=result["revisao"]; await db.commit()
    except Exception: await db.rollback(); raise
    result["total"]=operation.total_registros; logger.info("import_complete id=%s created=%s updated=%s errors=%s",operation.id,result["criados"],result["atualizados"],result["erros"]); return dict(result),created_ids
