from pathlib import Path
from unittest.mock import MagicMock

from docx import Document

from app.services.tabela_frete.analise import _analisar_documento_legacy
from app.services.tabela_frete.calculo_universal import calcular_universal
from app.services.tabela_frete.proposta_cif_docx import (
    extract_cif_proposal_docx,
    parse_cif_proposal_rows,
)


ROWS = [
    ["Destino", "Frete Peso", "Frete %", "Frete Mínimo", "Prazos de Entrega"],
    ["Araguaína – TO", "R$ 1,25", "7%", "R$ 250,00", "10 dias úteis"],
    ["Região", "R$ 1,25", "7%", "R$ 250,00", "10 – 15 dias úteis"],
    ["Balsas - MA", "R$ 1,25", "7%", "R$ 250,00", "10 dias úteis"],
]


def test_parser_normaliza_regras_calculaveis_sem_inventar_malha_regional():
    data = parse_cif_proposal_rows(
        ROWS,
        source_document="proposta.docx",
        full_text="Seguro 1%\nICMS: Conforme Legislação em Vigor\nCUBAGEM: 300K por M³",
    )

    assert data is not None
    assert data["formato"] == "tabela_frete_universal_v1"
    assert len(data["destinations"]) == 3
    araguaina = data["destinations"][0]
    assert (araguaina["city"], araguaina["uf"], araguaina["delivery_days"]) == ("ARAGUAINA", "TO", 10)
    assert araguaina["tariff_rule"]["excess_rate_per_kg"] == 1.25
    assert araguaina["freight_percentage"] == .07
    assert araguaina["minimum_freight"] == 250
    assert data["mapped_regions"][0]["reference_city"] == "ARAGUAINA"
    assert data["mapped_regions"][0]["delivery_days_max"] == 15
    assert "XAMBIOA" in data["mapped_regions"][0]["cities"]
    assert data["unresolved_regions"] == []
    assert data["surcharges"][0]["value"] == .01
    assert data["tax_rules"][0]["default_rate"] == .07


def test_regra_extraida_calcula_peso_percentual_minimo_e_seguro():
    data = parse_cif_proposal_rows(ROWS, source_document="proposta.docx", full_text="Seguro 1%")

    result = calcular_universal(data, {
        "origem_cidade": "Guarulhos",
        "origem_uf": "SP",
        "destino_cidade": "Araguaína",
        "destino_uf": "TO",
        "peso": 100,
        "valor_nf": 1000,
    })

    assert result["frete_base"] == 250
    assert result["valor_total"] == 279.57
    assert result["prazo_dias"] == 10


def test_barcarena_usa_automaticamente_a_regiao_de_belem():
    rows = [
        ROWS[0],
        ["Belém - PA", "R$ 1,25", "7%", "R$ 250,00", "15 dias úteis"],
        ["Região", "R$ 1,25", "7%", "R$ 250,00", "20 dias úteis"],
    ]
    data = parse_cif_proposal_rows(rows, source_document="proposta.docx", full_text="Seguro 1%")

    result = calcular_universal(data, {
        "origem_cidade": "Guarulhos", "origem_uf": "SP",
        "destino_cep": "68447000", "destino_cidade": "Barcarena", "destino_uf": "PA",
        "peso": .001, "valor_nf": 5000, "volume_total_m3": .37 * .57 * .57,
    })

    assert result["valor_total"] == 430.11
    assert result["prazo_dias"] == 20
    assert result["peso_considerado_kg"] == 36.064
    assert result["destino_tabela"]["regiao"] == "IBGE_IMEDIATA_150001"


def test_sao_domingos_do_maranhao_usa_regiao_comercial_de_bacabal():
    rows = [
        ROWS[0],
        ["Bacabal - MA", "R$ 1,25", "7%", "R$ 250,00", "10 dias úteis"],
        ["Região", "R$ 1,25", "7%", "R$ 250,00", "10 – 15 dias úteis"],
    ]
    data = parse_cif_proposal_rows(rows, source_document="proposta.docx", full_text="Seguro 1%")

    result = calcular_universal(data, {
        "origem_cep": "07042180", "origem_cidade": "Guarulhos", "origem_uf": "SP",
        "destino_cep": "65790000", "destino_cidade": "São Domingos do Maranhão",
        "destino_uf": "MA", "peso": 10, "valor_nf": 5000,
        "volume_total_m3": .57 * .57 * .34,
    })

    assert result["valor_total"] == 430.11
    assert result["prazo_dias"] == 15
    assert result["peso_considerado_kg"] == 33.14
    assert result["destino_tabela"]["regiao"] == "IBGE_IMEDIATA_210010"

    exact_city = calcular_universal(data, {
        "origem_cidade": "Guarulhos", "origem_uf": "SP",
        "destino_cep": "65700000", "destino_cidade": "Bacabal", "destino_uf": "MA",
        "peso": 10, "valor_nf": 5000, "volume_total_m3": .57 * .57 * .34,
    })
    assert exact_city["prazo_dias"] == 10
    assert exact_city["destino_tabela"]["cidade"] == "BACABAL"

    # Contratos analisados antes da inclusão das faixas continuam funcionando
    # pelo código regional que já estava persistido.
    for destination in data["destinations"]:
        destination.pop("cep_start", None)
        destination.pop("cep_end", None)
    legacy_result = calcular_universal(data, {
        "origem_cidade": "Guarulhos", "origem_uf": "SP",
        "destino_cep": "65790000", "destino_cidade": "São Domingos do Maranhão",
        "destino_uf": "MA", "peso": 10, "valor_nf": 5000,
        "volume_total_m3": .57 * .57 * .34,
    })
    assert legacy_result["valor_total"] == 430.11
    assert legacy_result["prazo_dias"] == 15


def test_docx_reconhecido_na_analise_em_vez_do_fallback_generico(tmp_path: Path):
    path = tmp_path / "proposta.docx"
    document = Document()
    document.add_paragraph("Seguro 1%")
    document.add_paragraph("ICMS: Conforme Legislação em Vigor")
    document.add_paragraph("CUBAGEM: 300K por M³")
    table = document.add_table(rows=0, cols=5)
    for values in ROWS:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
    document.save(path)

    parsed = extract_cif_proposal_docx(path)
    assert parsed is not None
    assert parsed["estatisticas"] == {
        "pracas": 3, "regras_tarifarias": 3, "regioes_mapeadas": 1,
        "regioes_sem_malha": 0,
    }

    source = MagicMock(nome_arquivo=path.name, caminho_storage=path.name, tipo_arquivo="docx")
    freight_table = MagicMock(transportadora_id="carrier-1")
    analysis = _analisar_documento_legacy(source, freight_table, tmp_path)

    assert analysis["dados_extraidos"]["formato"] == "tabela_frete_universal_v1"
    assert analysis["resumo"]["pracas"] == 3
    assert analysis["campos_com_duvida"] == []
