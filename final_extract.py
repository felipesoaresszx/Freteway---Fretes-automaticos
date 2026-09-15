#!/usr/bin/env python3
"""
Script final de extração de dados ALFA e RISPA.
Gera todos os arquivos JSON em data/tariffs/{alfa,rispa}/
"""

import json
import re
import os
import sys
from pathlib import Path
from typing import List, Dict, Any, Optional, Tuple
from datetime import datetime

try:
    import pymupdf
except ImportError:
    import fitz as pymupdf

try:
    import openpyxl
except ImportError:
    print("ERRO: openpyxl não está instalado")
    sys.exit(1)

# =============================================================================
# CONFIGURAÇÃO
# =============================================================================

BASE_DIR = Path(__file__).parent
FIXTURES_DIR = BASE_DIR / "backend" / "tests" / "fixtures"
OUTPUT_DIR = BASE_DIR / "data" / "tariffs"

ALFA_PDF = FIXTURES_DIR / "TABELA ALFA.pdf"
RISPA_XLSX = FIXTURES_DIR / "TABELA RISPA TODO BRASIL V1 26 (1).xlsx"

# Garantir diretórios
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
(OUTPUT_DIR / "alfa").mkdir(parents=True, exist_ok=True)
(OUTPUT_DIR / "rispa").mkdir(parents=True, exist_ok=True)


# =============================================================================
# FUNÇÕES AUXILIARES
# =============================================================================

def fix_encoding(text: str) -> str:
    """Corrige problemas de encoding do PDF."""
    if not text:
        return text
    encoding_map = {
        'Ò': 'Ô', 'Ï': 'Í', 'Ð': 'Ò', 'Ì': 'Í', 'þ': 'ç', 'Û': 'Ú',
        'Ù': 'Ú', 'Õ': 'Ô', '├': 'Ã', 'Þ': 'Ç', 'á': 'á', 'ß': 'ß',
        'Ý': 'Í', 'Ó': 'Ó', 'Ë': 'Ê', 'È': 'È', 'Î': 'Î', 'Å': 'Ã',
        'Ä': 'Ä', 'Ö': 'Ö', 'Ü': 'Ü', 'Ñ': 'Ñ', ' ': ' '
    }
    for wrong, correct in encoding_map.items():
        text = text.replace(wrong, correct)
    return text


def normalize_text(text: str) -> str:
    """Normaliza texto: uppercase, sem acentos, espaços colapsados."""
    if not text:
        return ""
    text = fix_encoding(text)
    text = re.sub(r'[\x00-\x1f\x7f-\x9f]', '', text)
    text = re.sub(r'[àáâãäå]', 'a', text)
    text = re.sub(r'[èéêë]', 'e', text)
    text = re.sub(r'[ìíîï]', 'i', text)
    text = re.sub(r'[òóôõö]', 'o', text)
    text = re.sub(r'[ùúûü]', 'u', text)
    text = re.sub(r'[ç]', 'c', text)
    text = re.sub(r'[ñ]', 'n', text)
    text = text.upper()
    text = re.sub(r'\s+', ' ', text)
    text = re.sub(r'[^A-Z0-9 ,.\-]', '', text)
    return text.strip()


def normalize_cep(cep: str) -> Optional[str]:
    """Normaliza CEP para string de 8 dígitos."""
    if not cep:
        return None
    cep_clean = re.sub(r'[^0-9]', '', str(cep))
    if len(cep_clean) != 8:
        return None
    return cep_clean


def parse_currency(value: Any) -> Optional[float]:
    """Parseia valor monetário para float."""
    if value is None:
        return None
    value = str(value).strip()
    value = re.sub(r'[R\$\.\s]', '', value)
    value = value.replace(',', '.')
    if not re.match(r'^[-+]?\d*\.?\d+$', value):
        return None
    try:
        return round(float(value), 6)
    except ValueError:
        return None


def parse_percentage_rispa(value: Any) -> Optional[float]:
    """Parseia percentual da RISPA (já é decimal: 0.006 = 0.6%)."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        try:
            return round(float(value), 6)
        except (ValueError, TypeError):
            return None
    value = str(value).strip()
    value = value.replace('%', '').replace(',', '.')
    try:
        return round(float(value) / 100, 6)
    except ValueError:
        return None


def parse_percentage_alfa(value: Any) -> Optional[float]:
    """Parseia percentual da ALFA (vém como texto: 0,40% = 0.40%)."""
    if value is None:
        return None
    value = str(value).strip()
    value = value.replace('%', '').replace(',', '.')
    try:
        return round(float(value) / 100, 6)
    except ValueError:
        return None


def parse_integer(value: Any) -> Optional[int]:
    """Parseia valor para inteiro."""
    if value is None:
        return None
    value = str(value).strip()
    try:
        return int(float(value))
    except (ValueError, TypeError):
        return None


def parse_date(date_str: str) -> Optional[str]:
    """Parseia data para formato ISO (YYYY-MM-DD)."""
    if not date_str:
        return None
    date_str = re.sub(r'[^0-9/]', '', date_str)
    try:
        if '/' in date_str:
            day, month, year = date_str.split('/')
            if len(year) == 2:
                year = f"20{year}"
            dt = datetime(int(year), int(month), int(day))
            return dt.strftime('%Y-%m-%d')
    except (ValueError, AttributeError):
        return None


def save_json(data: Any, file_path: Path):
    """Salva dados como JSON."""
    with open(file_path, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def generate_readme(pendencias: List[str], transportadora: str):
    """Gera README.md com pendências."""
    output_path = OUTPUT_DIR / transportadora / "README.md"
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(f"# Pendências - {transportadora}\n\n")
        f.write("Este documento lista todas as pendências identificadas durante a extração de dados.\n\n")
        f.write("## Itens a resolver\n\n")
        for i, item in enumerate(pendencias, 1):
            f.write(f"{i}. {item}\n\n")
        if not pendencias:
            f.write("Nenhuma pendência identificada.\n")


# =============================================================================
# EXTRAÇÃO ALFA
# =============================================================================

def extract_alfa():
    """Extrai todos os dados da ALFA."""
    print("Processando ALFA...")
    
    # --- Metadados ---
    doc = pymupdf.open(ALFA_PDF)
    full_text = ""
    for page in doc:
        full_text += page.get_text() + "\n"
    full_text = fix_encoding(full_text)
    
    metadata = {
        "transportadora": "ALFA TRANSPORTES",
        "cnpj": None,
        "cliente_tenant": "MODIAL COMERCIO DE ARTIGOS FUNERARIOS LTDA",
        "cliente_tenant_cnpj": None,
        "origem_uf": "SP",
        "origem_cidade": "GUARULHOS",
        "codigo_filial": None,
        "vigencia_inicio": None,
        "vigencia_fim": None,
        "versao_tabela": None,
        "cubagem_kg_m3": 250.0,
        "data_emissao": None
    }
    
    # CNPJ
    cnpj_match = re.search(r'(\d{2}\.?\d{3}\.?\d{3}/\d{4}-\d{2})', full_text)
    if cnpj_match:
        metadata["cliente_tenant_cnpj"] = re.sub(r'[./-]', '', cnpj_match.group(1))
    
    # Filial
    filial_match = re.search(r'Filial:\s*(\d+)', full_text, re.IGNORECASE)
    if filial_match:
        metadata["codigo_filial"] = filial_match.group(1)
    
    # Emissão
    emissao_match = re.search(r'Emiss[\S\s]*?(\d{2}/\d{2}/\d{4})', full_text, re.IGNORECASE)
    if emissao_match:
        metadata["data_emissao"] = parse_date(emissao_match.group(1))
    
    # Versão
    rev_match = re.search(r'Rev\.?\s*(\d+)', full_text, re.IGNORECASE)
    if rev_match:
        metadata["versao_tabela"] = f"REV {rev_match.group(1)}"
    
    save_json(metadata, OUTPUT_DIR / "alfa" / "metadata.json")
    
    # --- Destinos ---
    all_lines = []
    for page in doc:
        text = page.get_text()
        text = fix_encoding(text)
        all_lines.extend(text.split('\n'))
    
    # Encontrar início da tabela
    # Procurar por "Tarifa de pre" ou "TARIFA DE PRE" ou qualquer variação
    table_start = None
    for i, line in enumerate(all_lines):
        line_clean = line.strip()
        # Procurar por linha que contenha "Tarifa" ou "pre" (com ou sem acento)
        if ('ARIFA' in line_clean.upper() and 'PRE' in line_clean.upper()) or \
           ('ARIFA' in line_clean.upper() and any(x in line_clean.upper() for x in ['CO', 'ÇO', 'COB'])):
            table_start = i
            break
    
    # Se não encontrar, procurar por linha que tenha apenas "Origem"
    if table_start is None:
        for i, line in enumerate(all_lines):
            if line.strip().upper() in ['ORIGEM', 'DESTINO']:
                table_start = i - 1
                break
    
    if table_start is None:
        print("ERRO: Não foi possível encontrar o início da tabela ALFA")
        return
    
    table_lines = all_lines[table_start:]
    
    # Parsear entradas
    current_uf = None
    current_destino_lines = []
    current_codigo = None
    in_destino = False
    entries = []
    
    for line in table_lines:
        clean_line = line.strip()
        if not clean_line:
            continue
        
        # Cabeçalho
        if any(x in clean_line.upper() for x in ['ORIGEM', 'DESTINO', 'TARIFA', 'ATE 10', 
                                                   'TAXA EMBARQUE', 'FRETE VALOR', 'FRETE PESO']):
            in_destino = True
            continue
        
        if not in_destino:
            continue
        
        # UF
        uf_match = re.match(r'^[A-Z]{2}$', clean_line)
        if uf_match:
            if current_uf and current_destino_lines and current_codigo:
                entries.append({'uf': current_uf, 'destino': ' '.join(current_destino_lines), 'codigo': current_codigo})
            current_uf = clean_line
            current_destino_lines = []
            current_codigo = None
            in_destino = True
            continue
        
        # Código de tarifa
        codigo_match = re.match(r'^(\d{4})', clean_line)
        if codigo_match:
            current_codigo = codigo_match.group(1)
            in_destino = False
            continue
        
        # Cidade
        if in_destino:
            if clean_line.upper() not in ['ACIMA DE 100 KG', 'ACIMA DE 0', '%_R$', 'EXCEDENTE',
                                         'KG', 'ATE 10', 'ATE 30', 'ATE 50', 'ATE 70', 'ATE 100',
                                         'TAXA EMBARQUE', 'FRETE VALOR', 'FRETE PESO', 'KG (R$)']:
                current_destino_lines.append(clean_line)
            continue
    
    # Salvar última entrada
    if current_uf and current_destino_lines and current_codigo:
        entries.append({'uf': current_uf, 'destino': ' '.join(current_destino_lines), 'codigo': current_codigo})
    
    # Processar entradas para destinos
    special_zones = ["CAMPO GRANDE INTERIOR", "CAMPO GRANDE REGIÃO SUL", 
                     "CUIABÁ INTERIOR NORTÃO", "CUIABÁ - VALE DO ARAGUAIA",
                     "JAU (INTERIOR)", "IPAMERI (VIRTUAL)"]
    
    ufs_list = ['AC', 'AL', 'AP', 'AM', 'BA', 'CE', 'DF', 'ES', 'GO', 'MA', 'MT', 'MS',
                'MG', 'PA', 'PB', 'PR', 'PE', 'PI', 'RJ', 'RN', 'RS', 'RO', 'RR',
                'SC', 'SP', 'SE', 'TO']
    
    destinos = []
    zonas_data = {}
    
    for entry in entries:
        destino_original = entry['destino']
        
        # Separar destinos múltiplos
        if ',' in destino_original:
            destino_list = [d.strip() for d in destino_original.split(',')]
        else:
            # Verificar se há siglas de UF no meio
            words = destino_original.split()
            found_uf_in_middle = any(word in ufs_list for word in words[1:])
            
            if found_uf_in_middle:
                destino_list = []
                current = []
                for word in words:
                    if word in ufs_list:
                        if current:
                            destino_list.append(' '.join(current))
                            current = []
                    else:
                        current.append(word)
                if current:
                    destino_list.append(' '.join(current))
            else:
                destino_list = [destino_original]
        
        for destino_single in destino_list:
            eh_zona_especial = any(special.upper() in destino_single.upper() for special in special_zones)
            destinos.append({
                "uf": entry['uf'],
                "cidade_original": destino_single,
                "cidade_normalizada": normalize_text(destino_single) if not eh_zona_especial else None,
                "codigo_tarifa": entry['codigo'],
                "eh_zona_especial": eh_zona_especial
            })
        
        if entry['codigo'] not in zonas_data:
            zonas_data[entry['codigo']] = {
                "zona_id": entry['codigo'],
                "transportadora": "ALFA",
                "faixas": [],
                "excedente_kg": None,
                "taxa_embarque_minima": None
            }
    
    save_json(destinos, OUTPUT_DIR / "alfa" / "destinos.json")
    
    # --- Faixas de peso ---
    codigo_values = {}
    current_codigo = None
    
    for line in table_lines:
        clean_line = line.strip()
        if not clean_line:
            continue
        
        codigo_match = re.match(r'^(\d{4})\s+([\d.,]+)', clean_line)
        if codigo_match:
            current_codigo = codigo_match.group(1)
            if current_codigo not in codigo_values:
                codigo_values[current_codigo] = []
            codigo_values[current_codigo].append(parse_currency(codigo_match.group(2)))
            continue
        
        if current_codigo and re.match(r'^[\d.,]+$', clean_line):
            codigo_values[current_codigo].append(parse_currency(clean_line))
            continue
        
        if current_codigo and not re.match(r'^[\d.,]+$', clean_line) and not re.match(r'^\d{4}', clean_line):
            current_codigo = None
    
    for codigo, values in codigo_values.items():
        if codigo in zonas_data:
            weights = [10, 30, 50, 70, 100]
            for i, weight in enumerate(weights):
                if i < len(values) and values[i] is not None:
                    zonas_data[codigo]["faixas"].append({"peso_ate_kg": weight, "valor": values[i]})
            if len(values) > 5:
                zonas_data[codigo]["excedente_kg"] = values[5]
    
    save_json(list(zonas_data.values()), OUTPUT_DIR / "alfa" / "faixas_peso.json")
    
    # --- Percentuais (do prompt) ---
    percentuais = [
        {"zona_id": "*", "componente": "frete_valor", "percentual": 0.0040, "minimo": None, "excecao_uf": None},
        {"zona_id": "*", "componente": "gris", "percentual": 0.0015, "minimo": 6.75, "excecao_uf": None},
        {"zona_id": "*", "componente": "gris", "percentual": 0.0060, "minimo": 9.18, "excecao_uf": "RJ"}
    ]
    save_json(percentuais, OUTPUT_DIR / "alfa" / "percentuais.json")
    
    # --- Arquivos vazios ---
    save_json([], OUTPUT_DIR / "alfa" / "taxas_fixas.json")
    save_json({"transportadora": "ALFA", "taxas_globais": []}, OUTPUT_DIR / "alfa" / "taxas_globais.json")
    save_json([], OUTPUT_DIR / "alfa" / "adicionais.json")
    save_json([], OUTPUT_DIR / "alfa" / "prazos.json")
    
    # --- README ---
    pendencias = [
        "Taxas fixas por zona: não identificadas no PDF - necessário verificar documento original",
        "Taxas globais: não identificadas no PDF - necessário verificar seção 'Generalidades e Serviços Adicionais'",
        "Adicionais condicionais: não extraídos do PDF - necessário verificar documento",
        "Prazo de entrega: NÃO EXISTE no PDF - todos os campos são null"
    ]
    generate_readme(pendencias, "alfa")
    
    print(f"  Metadados: {len(metadata)} campos")
    print(f"  Destinos: {len(destinos)} registros")
    print(f"  Faixas de peso: {len(zonas_data)} zonas")
    print(f"  Percentuais: {len(percentuais)} registros")


# =============================================================================
# EXTRAÇÃO RISPA
# =============================================================================

def extract_rispa():
    """Extrai todos os dados da RISPA."""
    print("Processando RISPA...")
    
    wb = openpyxl.load_workbook(RISPA_XLSX, data_only=True)
    
    # --- Metadados ---
    metadata = {
        "transportadora": "RISPA TRANSPORTES",
        "cnpj": None,
        "cliente_tenant": None,
        "cliente_tenant_cnpj": None,
        "origem_uf": "SP",
        "origem_cidade": "GUARULHOS",
        "codigo_filial": None,
        "vigencia_inicio": None,
        "vigencia_fim": None,
        "versao_tabela": None,
        "cubagem_kg_m3": 300.0,
        "data_emissao": None
    }
    
    sheet_proposta = wb["PROPOSTA TAB FRETE RISPA"]
    for row in sheet_proposta.iter_rows(values_only=True):
        row_text = " ".join([str(cell) if cell else "" for cell in row]).upper()
        
        if "RISPA" in row_text:
            metadata["transportadora"] = "RISPA TRANSPORTES"
        
        if "VERSÃO" in row_text or "VERSAO" in row_text:
            versao_match = re.search(r'VERS[\S\s]*?(\d+\.?\d*)', row_text)
            if versao_match:
                metadata["versao_tabela"] = versao_match.group(1)
        
        cnpj_match = re.search(r'(\d{2}\.?\d{3}\.?\d{3}/\d{4}-\d{2})', row_text)
        if cnpj_match:
            metadata["cliente_tenant_cnpj"] = re.sub(r'[./-]', '', cnpj_match.group(1))
            metadata["cliente_tenant"] = "MODIAL COMERCIO DE ARTIGOS FUNERARIOS LTDA"
        
        if "GUARULHOS" in row_text:
            metadata["origem_cidade"] = "GUARULHOS"
            metadata["origem_uf"] = "SP"
        
        date_match = re.search(r'(\d{2}/\d{2}/\d{4})', row_text)
        if date_match:
            if not metadata.get("data_emissao"):
                metadata["data_emissao"] = parse_date(date_match.group(1))
    
    # Cubagem
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        for row in sheet.iter_rows(values_only=True):
            row_text = " ".join([str(cell) if cell else "" for cell in row]).upper()
            cubagem_match = re.search(r'CUBAGEM[\S\s]*?(\d+)\s*KG\s*m\s*3', row_text, re.IGNORECASE)
            if cubagem_match:
                metadata["cubagem_kg_m3"] = float(cubagem_match.group(1))
                break
        if metadata.get("cubagem_kg_m3"):
            break
    
    save_json(metadata, OUTPUT_DIR / "rispa" / "metadata.json")
    
    # --- Resolução CEP ---
    sheet_cep = wb["TABELA FAIXA CEP"]
    header = [cell.value for cell in sheet_cep[1]]
    
    # Encontrar índices
    uf_idx = None
    cidade_idx = None
    sigla_idx = None
    cep_i_idx = None
    cep_f_idx = None
    prazo_idx = None
    tda_idx = None
    
    for idx, col in enumerate(header):
        col_upper = str(col).upper().strip() if col else ""
        if col_upper == "UF":
            uf_idx = idx
        elif col_upper == "CIDADE":
            cidade_idx = idx
        elif col_upper == "SIGLA":
            sigla_idx = idx
        elif col_upper in ["CEPI", "CEP I"]:
            cep_i_idx = idx
        elif col_upper in ["CEPF", "CEP F"]:
            cep_f_idx = idx
        elif "PRAZO" in col_upper:
            prazo_idx = idx
        elif col_upper == "TDA":
            tda_idx = idx
    
    cep_faixas = []
    errors = []
    
    for row in sheet_cep.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        
        uf = str(row[uf_idx]).strip() if uf_idx is not None and uf_idx < len(row) else None
        cidade = str(row[cidade_idx]).strip() if cidade_idx is not None and cidade_idx < len(row) else None
        sigla_zona = str(row[sigla_idx]).strip() if sigla_idx is not None and sigla_idx < len(row) else None
        
        cep_inicio = normalize_cep(str(row[cep_i_idx])) if cep_i_idx is not None and cep_i_idx < len(row) else None
        cep_fim = normalize_cep(str(row[cep_f_idx])) if cep_f_idx is not None and cep_f_idx < len(row) else None
        
        prazo = parse_integer(str(row[prazo_idx])) if prazo_idx is not None and prazo_idx < len(row) else None
        tda = parse_currency(str(row[tda_idx])) if tda_idx is not None and tda_idx < len(row) else None
        
        if prazo == 0:
            prazo = None
        
        # Validar CEP
        if cep_inicio and len(cep_inicio) != 8:
            errors.append(f"CEP inválido: {row[cep_i_idx] if cep_i_idx else 'N/A'} ({uf}/{cidade})")
            continue
        
        if cep_fim and len(cep_fim) != 8:
            errors.append(f"CEP inválido: {row[cep_f_idx] if cep_f_idx else 'N/A'} ({uf}/{cidade})")
            continue
        
        cep_faixas.append({
            "uf": uf,
            "cidade": cidade,
            "sigla_zona": sigla_zona,
            "cep_inicio": cep_inicio,
            "cep_fim": cep_fim,
            "prazo_dias_util": prazo,
            "tda": tda
        })
    
    save_json(cep_faixas, OUTPUT_DIR / "rispa" / "resolucao_destino_cep.json")
    
    # --- Cidades atendidas ---
    cidades_atendidas = []
    if "CIDADES ATENDIDAS" in wb.sheetnames:
        sheet_cidades = wb["CIDADES ATENDIDAS"]
        header = [cell.value for cell in sheet_cidades[1]]
        
        uf_idx = None
        cidade_idx = None
        sigla_idx = None
        prazo_idx = None
        tda_idx = None
        
        for idx, col in enumerate(header):
            col_upper = str(col).upper().strip() if col else ""
            if col_upper == "UF":
                uf_idx = idx
            elif col_upper == "CIDADE DESTINO":
                cidade_idx = idx
            elif col_upper == "SIGLA":
                sigla_idx = idx
            elif col_upper == "PRAZO":
                prazo_idx = idx
            elif col_upper == "TDA":
                tda_idx = idx
        
        for row in sheet_cidades.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue
            
            uf = str(row[uf_idx]).strip() if uf_idx is not None and uf_idx < len(row) else None
            cidade = str(row[cidade_idx]).strip() if cidade_idx is not None and cidade_idx < len(row) else None
            sigla_zona = str(row[sigla_idx]).strip() if sigla_idx is not None and sigla_idx < len(row) else None
            prazo = parse_integer(str(row[prazo_idx])) if prazo_idx is not None and prazo_idx < len(row) else None
            tda = parse_currency(str(row[tda_idx])) if tda_idx is not None and tda_idx < len(row) else None
            
            if prazo == 0:
                prazo = None
            
            cidades_atendidas.append({
                "uf": uf,
                "cidade": cidade,
                "sigla_zona": sigla_zona,
                "prazo_dias": prazo,
                "tda": tda
            })
    
    save_json(cidades_atendidas, OUTPUT_DIR / "rispa" / "resolucao_destino_cidades.json")
    
    # --- Faixas de peso ---
    sheet_cep = wb["TABELA FAIXA CEP"]
    header = [cell.value for cell in sheet_cep[1]]
    
    sigla_idx = None
    for idx, col in enumerate(header):
        col_upper = str(col).upper().strip() if col else ""
        if col_upper == "SIGLA":
            sigla_idx = idx
    
    # Encontrar colunas de peso
    weight_cols = {}
    for idx, col_name in enumerate(header):
        if col_name and "KG" in str(col_name).upper() and "AT" in str(col_name).upper():
            weight_match = re.search(r'AT[ÉE]\s+(\d+)', str(col_name), re.IGNORECASE)
            if weight_match:
                weight_cols[int(weight_match.group(1))] = idx
    
    # Encontrar coluna de excedente
    excedente_idx = None
    for idx, col_name in enumerate(header):
        if col_name and "EXCEDENTE" in str(col_name).upper():
            excedente_idx = idx
            break
    
    zonas_data = {}
    
    for row in sheet_cep.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        
        sigla_zona = str(row[sigla_idx]).strip() if sigla_idx is not None and sigla_idx < len(row) else None
        
        if not sigla_zona:
            continue
        
        if sigla_zona not in zonas_data:
            zonas_data[sigla_zona] = {
                "zona_id": sigla_zona,
                "transportadora": "RISPA",
                "faixas": [],
                "excedente_kg": None,
                "taxa_embarque_minima": None
            }
        
        # Extrair faixas
        for peso, col_idx in weight_cols.items():
            if col_idx < len(row):
                valor = parse_currency(row[col_idx])
                if valor is not None:
                    zonas_data[sigla_zona]["faixas"].append({
                        "peso_ate_kg": peso,
                        "valor": valor,
                        "valor_bruto": row[col_idx]
                    })
        
        # Extrair excedente
        if excedente_idx and excedente_idx < len(row):
            valor = parse_currency(row[excedente_idx])
            if valor is not None:
                zonas_data[sigla_zona]["excedente_kg"] = valor
    
    save_json(list(zonas_data.values()), OUTPUT_DIR / "rispa" / "faixas_peso.json")
    
    # --- Percentuais ---
    gris_idx = None
    for idx, col_name in enumerate(header):
        if col_name and "GRIS" in str(col_name).upper() and "ADV" in str(col_name).upper():
            gris_idx = idx
            break
    
    zonas_percentuais = {}
    if gris_idx:
        for row in sheet_cep.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue
            
            sigla_zona = str(row[sigla_idx]).strip() if sigla_idx is not None and sigla_idx < len(row) else None
            gris_valor = row[gris_idx] if gris_idx < len(row) else None
            
            if sigla_zona and gris_valor and sigla_zona not in zonas_percentuais:
                percentual = parse_percentage_rispa(gris_valor)
                if percentual is not None:
                    zonas_percentuais[sigla_zona] = {
                        "zona_id": sigla_zona,
                        "componente": "gris_advalorem",
                        "percentual": percentual,
                        "minimo": 6.70,
                        "excecao_uf": None
                    }
    
    save_json(list(zonas_percentuais.values()), OUTPUT_DIR / "rispa" / "percentuais.json")
    
    # --- Taxas fixas ---
    col_indices = {}
    for idx, col_name in enumerate(header):
        if col_name:
            col_upper = str(col_name).upper().strip()
            if col_upper in ["TAS", "TRT", "PEDÁGIO", "PEDAGIO", "TAXA DESPACHO", "TDA"]:
                col_indices[col_upper] = idx
    
    taxas_dict = {}
    for row in sheet_cep.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        
        sigla_zona = str(row[sigla_idx]).strip() if sigla_idx is not None and sigla_idx < len(row) else None
        
        if not sigla_zona:
            continue
        
        if sigla_zona not in taxas_dict:
            taxas_dict[sigla_zona] = {
                "zona_id": sigla_zona,
                "tas": None, "trt": None, "trt_regra_texto": "",
                "pedagio_fracao_100kg": None, "taxa_despacho": None, "tda": None
            }
        
        for col_name, col_idx in col_indices.items():
            if col_idx < len(row):
                valor = parse_currency(row[col_idx])
                if col_name == "TAS":
                    taxas_dict[sigla_zona]["tas"] = valor
                elif col_name == "TRT":
                    taxas_dict[sigla_zona]["trt"] = valor
                elif col_name in ["PEDÁGIO", "PEDAGIO"]:
                    taxas_dict[sigla_zona]["pedagio_fracao_100kg"] = valor
                elif col_name == "TAXA DESPACHO":
                    taxas_dict[sigla_zona]["taxa_despacho"] = valor
                elif col_name == "TDA":
                    taxas_dict[sigla_zona]["tda"] = valor
    
    save_json(list(taxas_dict.values()), OUTPUT_DIR / "rispa" / "taxas_fixas.json")
    
    # --- Taxas globais ---
    taxas_globais = []
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        for row in sheet.iter_rows(values_only=True):
            row_text = " ".join([str(cell) if cell else "" for cell in row]).upper()
            for codigo, base in [("TAD", "por CTe"), ("TEC", "por CTe"), ("TPF", "por entrega, destinatário PF")]:
                if codigo in row_text and re.search(r'R\$\s*[\d.,]+', row_text):
                    valor_match = re.search(r'R\$\s*([\d.,]+)', row_text)
                    if valor_match:
                        taxas_globais.append({"codigo": codigo, "valor": parse_currency(valor_match.group(1)), "base": base})
    
    save_json({"transportadora": "RISPA", "taxas_globais": taxas_globais}, OUTPUT_DIR / "rispa" / "taxas_globais.json")
    
    # --- Adicionais ---
    adicionais = []
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        for row in sheet.iter_rows(values_only=True):
            row_text = " ".join([str(cell) if cell else "" for cell in row])
            for codigo, desc in [("TDE", "Taxa de Descarregamento"), ("TZR", "Taxa de Zona Rural"),
                                  ("PALETIZACAO", "Paletização"), ("ARMAZENAGEM", "Armazenagem"),
                                  ("DEVOLUCAO", "Devolução"), ("REENTREGA", "Reentrega")]:
                if codigo in row_text.upper():
                    adicionais.append({"codigo": codigo, "descricao": desc, "valor_ou_regra": None, "condicao": row_text})
            if "PRODUTO QUÍMICO" in row_text.upper() or "CLASSE" in row_text.upper():
                adicionais.append({"codigo": "PRODUTO_QUIMICO", "descricao": "Produto químico", "valor_ou_regra": None, "condicao": row_text})
    
    if "ONDE COLETAMOS" in wb.sheetnames:
        sheet = wb["ONDE COLETAMOS"]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue
            uf = str(row[0]).strip() if len(row) > 0 else None
            cidade = str(row[1]).strip() if len(row) > 1 else None
            distancia_km = parse_currency(row[2]) if len(row) > 2 else None
            taxa_coleta = parse_currency(row[3]) if len(row) > 3 else None
            if uf and cidade:
                adicionais.append({"codigo": "COLETA_FORA_GUARULHOS", "descricao": "Coleta fora de Guarulhos", "valor_ou_regra": taxa_coleta, "condicao": f"{uf}/{cidade}, distancia: {distancia_km} km"})
    
    save_json(adicionais, OUTPUT_DIR / "rispa" / "adicionais.json")
    
    # --- Prazos ---
    prazos = []
    for sheet_name in ["PROPOSTA TAB FRETE RISPA", "TABELA FAIXA CEP"]:
        if sheet_name not in wb.sheetnames:
            continue
        sheet = wb[sheet_name]
        for row in sheet.iter_rows(values_only=True):
            row_text = " ".join([str(cell) if cell else "" for cell in row]).upper()
            prazo_match = re.search(r'(\d+)\s+[Aa]\s+(\d+)\s+DIAS\s+UTEIS', row_text, re.IGNORECASE)
            if prazo_match:
                min_dias = int(prazo_match.group(1))
                max_dias = int(prazo_match.group(2))
                for cell in row:
                    if cell and re.match(r'^[A-Z]+$', str(cell).upper()):
                        sigla_zona = str(cell).upper()
                        prazos.append({"zona_id": sigla_zona, "prazo_min_dias": min_dias, "prazo_max_dias": max_dias, "fonte": sheet_name})
                        break
    
    save_json(prazos, OUTPUT_DIR / "rispa" / "prazos.json")
    
    # --- README ---
    generate_readme(errors, "rispa")
    
    print(f"  Metadados: {len(metadata)} campos")
    print(f"  Faixas CEP: {len(cep_faixas)} registros")
    print(f"  Cidades atendidas: {len(cidades_atendidas)} registros")
    print(f"  Faixas de peso: {len(zonas_data)} zonas")
    print(f"  Percentuais: {len(zonas_percentuais)} registros")
    print(f"  Pendências: {len(errors)} itens")


# =============================================================================
# EXECUÇÃO
# =============================================================================

if __name__ == "__main__":
    print("=" * 60)
    print("Extração de dados ALFA e RISPA - v3")
    print("=" * 60)
    
    if not ALFA_PDF.exists():
        print(f"ERRO: Arquivo não encontrado: {ALFA_PDF}")
        sys.exit(1)
    
    if not RISPA_XLSX.exists():
        print(f"ERRO: Arquivo não encontrado: {RISPA_XLSX}")
        sys.exit(1)
    
    extract_alfa()
    print()
    extract_rispa()
    
    print("\n" + "=" * 60)
    print("Extração concluída!")
    print(f"Arquivos gerados em: {OUTPUT_DIR.absolute()}")
    print("=" * 60)
